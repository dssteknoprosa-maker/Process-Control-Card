"""
DSS Kebutuhan Bahan Kimia Pengolahan Kulit (versi Streamlit)
Database: Supabase Postgres. Deploy: Streamlit Community Cloud.

Secrets yang dibutuhkan:
  DEV_ID, DEV_PASS, USER_ID, USER_PASS, SUPABASE_URL, SUPABASE_KEY
"""
import hmac
import io
import os
import re
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal
from itertools import groupby

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from supabase import create_client

TABEL = "formulasi"
KOLOM = ["id", "jenis_kulit", "tahapan", "senyawa", "persen", "satuan", "fase"]
# Faktor konversi dari gram (berat kulit dalam g -> jumlah senyawa dalam g).
# mL dan L memakai asumsi massa jenis 1 g = 1 mL.
SATUAN = {"g": 1, "kg": 0.001, "mL": 1, "L": 0.001}
BARU = "➕ Tambah baru..."

st.set_page_config(page_title="DSS Kimia Pengolahan Kulit", page_icon="🧪", layout="wide")


# ---------------------------------------------------------------------------
# Konfigurasi & database
# ---------------------------------------------------------------------------
def cfg(nama: str) -> str:
    """Baca dari st.secrets, fallback ke environment variable."""
    try:
        if nama in st.secrets:
            return str(st.secrets[nama])
    except Exception:  # noqa: BLE001
        pass
    return os.getenv(nama, "")


class DBError(Exception):
    """Error database dengan pesan ramah pengguna."""


@st.cache_resource
def db():
    url, key = cfg("SUPABASE_URL"), cfg("SUPABASE_KEY")
    if not url or not key:
        raise DBError("SUPABASE_URL / SUPABASE_KEY belum diisi di Secrets.")
    try:
        return create_client(url, key)
    except Exception as e:  # noqa: BLE001
        raise DBError(f"Gagal membuat koneksi ke Supabase: {str(e)[:200]}")


def jalankan(fn):
    try:
        return fn()
    except DBError:
        raise
    except Exception as e:  # noqa: BLE001
        teks = str(e)
        if "23505" in teks or "duplicate" in teks.lower():
            raise DBError("Data duplikat: kombinasi jenis kulit + tahapan + senyawa sudah ada.")
        if ("satuan" in teks or "fase" in teks) and ("column" in teks or "PGRST204" in teks):
            raise DBError("Kolom baru (satuan/fase) belum ada. Jalankan migrasi SQL terbaru di Supabase SQL Editor.")
        if "does not exist" in teks or "PGRST205" in teks:
            raise DBError("Tabel database belum ada. Jalankan setup.sql dan semua file migrasi di Supabase SQL Editor.")
        raise DBError(f"Gagal mengakses database: {teks[:200]}")


def ambil_semua() -> pd.DataFrame:
    res = jalankan(
        lambda: db().table(TABEL).select("*").order("jenis_kulit").order("id").execute()
    )
    df = pd.DataFrame(res.data or [])
    if df.empty:
        return pd.DataFrame(columns=KOLOM)
    df["persen"] = df["persen"].astype(float)
    if "satuan" not in df:
        df["satuan"] = "kg"
    df["satuan"] = df["satuan"].fillna("kg")
    if "fase" not in df:
        df["fase"] = 1
    df["fase"] = pd.to_numeric(df["fase"], errors="coerce").fillna(1).astype(int)
    return df[KOLOM]


def ambil_info() -> dict:
    """Waktu & keterangan per (jenis kulit, tahapan, fase)."""
    res = jalankan(lambda: db().table("tahapan_info").select("*").execute())
    return {(r["jenis_kulit"], r["tahapan"], int(r.get("fase") or 1)):
            (r.get("waktu") or "", r.get("keterangan") or "")
            for r in (res.data or [])}


def rapikan(teks) -> str:
    return re.sub(r"\s+", " ", (teks or "").strip())


def flash(jenis: str, pesan: str):
    st.session_state["flash"] = (jenis, pesan)
    st.rerun()


def tampil_flash():
    if "flash" in st.session_state:
        jenis, pesan = st.session_state.pop("flash")
        getattr(st, jenis)(pesan)


# ---------------------------------------------------------------------------
# Login
# ---------------------------------------------------------------------------
def cocok(a: str, b: str, benar_a: str, benar_b: str) -> bool:
    if not benar_a or not benar_b:
        return False
    return (hmac.compare_digest(a.encode(), benar_a.encode())
            and hmac.compare_digest(b.encode(), benar_b.encode()))


def halaman_login():
    st.title("🧪 DSS Kebutuhan Bahan Kimia Pengolahan Kulit")
    with st.form("login"):
        uid = st.text_input("ID")
        pwd = st.text_input("Password", type="password")
        masuk = st.form_submit_button("Masuk", type="primary")
    if masuk:
        uid = uid.strip()
        if cocok(uid, pwd, cfg("DEV_ID"), cfg("DEV_PASS")):
            st.session_state["role"] = "dev"
            st.rerun()
        elif cocok(uid, pwd, cfg("USER_ID"), cfg("USER_PASS")):
            st.session_state["role"] = "user"
            st.rerun()
        else:
            st.error("❌ ID atau password salah.")


def logout_sidebar():
    with st.sidebar:
        st.write(f"Masuk sebagai: **{'Developer' if st.session_state['role'] == 'dev' else 'User'}**")
        if st.button("Logout"):
            st.session_state.clear()
            st.rerun()


# ---------------------------------------------------------------------------
# Interface Developer
# ---------------------------------------------------------------------------
def pilih_atau_baru(label: str, pilihan: list, kunci: str) -> str:
    """Selectbox dengan opsi tambah nilai baru."""
    terpilih = st.selectbox(label, pilihan + [BARU], key=f"sb_{kunci}")
    if terpilih == BARU:
        return st.text_input(f"{label} baru", key=f"tx_{kunci}")
    return terpilih


BARU_TAHAP = "➕ Tahapan baru tanpa bahan kimia"
MODE_LEWATI = "Lewati yang sudah ada (aman)"
MODE_TIMPA = "Timpa yang sudah ada (perbarui persen, satuan, fase, waktu, keterangan)"


def baca_formulasi(f, sheet=None) -> dict:
    """Ubah file Kartu Kendali Proses menjadi data formulasi + info waktu/keterangan per fase."""
    daftar, info = baca_file(f, sheet)
    rec, info_fase, catatan, langkah, tanpa = [], {}, [], set(), []
    for e in daftar:
        jenis, tahap = rapikan(e["jenis"]).title(), rapikan(e["tahapan"]).title()
        if (jenis, tahap) in langkah:
            catatan.append(f"{jenis} / {tahap} muncul lebih dari sekali di file; yang pertama dipakai.")
            continue
        langkah.add((jenis, tahap))
        if not e["rows"]:  # tahapan tanpa bahan kimia
            tb = e.get("tanpa_bahan", {})
            w, k = tb.get("waktu", ""), tb.get("keterangan", "")
            tanpa.append({"jenis_kulit": jenis, "tahapan": tahap, "waktu": w, "keterangan": k})
            info_fase[(jenis, tahap, 1)] = (w, k)
            continue
        nama_ada = set()
        for r in e["rows"]:
            nama, persen = rapikan(r["senyawa"]), r["persen"]
            if persen is None or pd.isna(persen) or persen <= 0:
                catatan.append(f"{jenis} / {tahap} / {nama}: persen tidak valid, dilewati.")
                continue
            if nama.lower() in nama_ada:
                catatan.append(f"{jenis} / {tahap} / {nama}: nama senyawa ganda, dilewati.")
                continue
            nama_ada.add(nama.lower())
            fase = int(r["fase"])
            rec.append({"jenis_kulit": jenis, "tahapan": tahap, "senyawa": nama,
                        "persen": float(persen), "satuan": r["satuan"], "fase": fase})
            w, k = (r.get("waktu") or "").strip(), (r.get("keterangan") or "").strip()
            if w or k:
                info_fase[(jenis, tahap, fase)] = (w, k)
    if info.get("gagal"):
        catatan.append(f"{info['gagal']} baris bahan kimia tidak terbaca "
                       "(formatnya harus 'persen (jumlah satuan)').")
    return {"formulasi": rec, "info": info_fase, "catatan": catatan, "tanpa_bahan": tanpa}


def simpan_impor(hasil: dict, mode: str) -> str:
    """Simpan hasil baca file ke database. Return ringkasan hasil."""
    ada = {(r.jenis_kulit, r.tahapan, r.senyawa.lower()): int(r.id)
           for r in ambil_semua().itertuples()}
    baru, ubah, lewat = [], [], 0
    for r in hasil["formulasi"]:
        kunci = (r["jenis_kulit"], r["tahapan"], r["senyawa"].lower())
        if kunci not in ada:
            baru.append(r)
        elif mode == MODE_TIMPA:
            ubah.append((ada[kunci], {"persen": r["persen"], "satuan": r["satuan"],
                                      "fase": r["fase"]}))
        else:
            lewat += 1
    if baru:
        jalankan(lambda: db().table(TABEL).insert(baru).execute())
    for id_, data in ubah:
        jalankan(lambda: db().table(TABEL).update(data).eq("id", id_).execute())

    n_info = 0
    try:
        info_ada = ambil_info()
        daftar_info = [{"jenis_kulit": j, "tahapan": t, "fase": f, "waktu": w, "keterangan": k}
                       for (j, t, f), (w, k) in hasil["info"].items()
                       if mode == MODE_TIMPA or (j, t, f) not in info_ada]
        if daftar_info:
            jalankan(lambda: db().table("tahapan_info").upsert(
                daftar_info, on_conflict="jenis_kulit,tahapan,fase").execute())
        n_info = len(daftar_info)
    except DBError as e:
        return (f"⚠️ Formulasi tersimpan ({len(baru)} baru, {len(ubah)} diperbarui, "
                f"{lewat} dilewati), tapi waktu & keterangan gagal: {e}")
    return (f"✅ Impor selesai: {len(baru)} bahan baru, {len(ubah)} diperbarui, "
            f"{lewat} dilewati; waktu & keterangan {n_info} fase/tahapan "
            f"({len(hasil.get('tanpa_bahan', []))} tahapan tanpa bahan kimia di file).")


def deteksi_format(f) -> dict:
    """Return {nama_sheet: 'kartu' | 'resep'} untuk sheet yang formatnya dikenali."""
    wb = load_workbook(f, data_only=True)
    hasil = {}
    for ws in wb.worksheets:
        for r in range(1, min(ws.max_row, 40) + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if isinstance(a, str) and isinstance(b, str):
                ka, kb = a.strip().lower(), b.strip().lower()
                if ka == "tanggal" and kb == "proses":
                    hasil[ws.title] = "kartu"
                    break
                if ka == "tahap" and kb == "bahan":
                    hasil[ws.title] = "resep"
                    break
    return hasil


def _teks(v) -> str:
    return "" if v is None else str(v).strip()


def _bentuk_fase(rows: list) -> list:
    """Kelompokkan bahan dalam satu tahap menjadi fase.
    Sel Waktu yang digabung = satu fase. Baris tanpa waktu & keterangan ikut fase
    berikutnya yang punya waktu (waktu biasanya ditulis di baris terakhir kelompok)."""
    unit = []
    for row in rows:
        if row["gid"] is not None and unit and unit[-1]["gid"] == row["gid"]:
            unit[-1]["rows"].append(row)
        else:
            unit.append({"gid": row["gid"], "rows": [row]})
    fases, cur = [], []
    for u in unit:
        cur += u["rows"]
        if any(nilai_teks(x["waktu"]) or nilai_teks(x["ket"]) for x in u["rows"]):
            fases.append(cur)
            cur = []
    if cur:
        if fases:
            fases[-1] += cur
        else:
            fases.append(cur)
    return fases


def _bersih(t: str) -> str:
    return (t or "").strip().rstrip(" ,;")


def baca_resep(f, sheet: str, jenis: str) -> dict:
    """Baca format resep: kolom Tahap | Bahan | Jumlah (%) | Waktu (menit) | Keterangan.
    Baris tanpa bahan kimia (mis. Aging, Shaving, Putaran dihentikan) dibaca sebagai
    tahapan tanpa bahan: hanya waktu & keterangan."""
    wb = load_workbook(f, data_only=True)
    ws = wb[sheet]

    kol, mulai = {}, None
    for r in range(1, min(ws.max_row, 40) + 1):
        a, b = _teks(ws.cell(r, 1).value).lower(), _teks(ws.cell(r, 2).value).lower()
        if a == "tahap" and b == "bahan":
            for k in range(1, 12):
                t = _teks(ws.cell(r, k).value).lower()
                for nama in ("tahap", "bahan", "jumlah", "waktu", "keterangan"):
                    if t.startswith(nama) and nama not in kol:
                        kol[nama] = k
            mulai = r + 1
            break
    if mulai is None or not all(k in kol for k in ("tahap", "bahan", "jumlah")):
        raise ValueError("Header Tahap | Bahan | Jumlah (%) tidak ditemukan di sheet ini.")
    kd, ke = kol.get("waktu"), kol.get("keterangan")

    grup_d = {}  # baris -> baris pertama dari sel Waktu yang digabung vertikal
    if kd:
        for rng in ws.merged_cells.ranges:
            if rng.max_row > rng.min_row and rng.min_col <= kd <= rng.max_col:
                for rr in range(rng.min_row, rng.max_row + 1):
                    grup_d[rr] = rng.min_row

    baris = []  # hanya baris yang tidak kosong
    for r in range(mulai, ws.max_row + 1):
        a = _teks(ws.cell(r, kol["tahap"]).value)
        b = _teks(ws.cell(r, kol["bahan"]).value)
        c = ws.cell(r, kol["jumlah"]).value
        d = ws.cell(grup_d.get(r, r), kd).value if kd else None
        e = ws.cell(r, ke).value if ke else None
        if a or b or c is not None or _teks(d) or _teks(e):
            baris.append({"r": r, "a": a, "b": b, "c": c, "d": d, "e": e})

    tahap_raw, tahap_final, terakhir = None, None, None
    kemunculan, data, seen, tanpa, lanjutan = {}, {}, {}, {}, {}
    catatan, kaki, diganti = [], {}, []

    def nama_unik(raw: str) -> str:
        kunci = raw.casefold()
        kemunculan[kunci] = kemunculan.get(kunci, 0) + 1
        dasar = rapikan(raw).title()
        if kemunculan[kunci] == 1:
            return dasar
        akhir = f"{dasar} ({kemunculan[kunci]})"
        diganti.append(f"Tahap '{dasar}' muncul lagi, dinamai '{akhir}'.")
        return akhir

    for i, x in enumerate(baris):
        r, a, b, c, d, e = (x[k] for k in ("r", "a", "b", "c", "d", "e"))
        m = re.match(r"^(?:keterangan\s*:)?\s*(\*+)\s*(.+)$", a, re.I)
        if m and not b:  # catatan kaki, mis. "Keterangan : ** dilarutkan ..."
            kaki[m.group(1)] = m.group(2).strip()
            continue

        if b:  # ---- baris bahan kimia ----
            if a and a.casefold() != (tahap_raw or "").casefold():
                tahap_raw, tahap_final = a, None
            try:
                persen = float(_teks(c).replace(",", "."))
            except ValueError:
                persen = None
            if persen is None or persen <= 0:
                catatan.append(f"Baris {r}: {b} dilewati (jumlah % tidak valid: {_teks(c) or 'kosong'}).")
                continue
            if tahap_raw is None:
                catatan.append(f"Baris {r}: {b} dilewati (belum ada nama tahap).")
                continue
            if tahap_final is None:
                tahap_final = nama_unik(tahap_raw)
                data[tahap_final], seen[tahap_final] = [], {}
            bintang = re.search(r"\*+$", b)
            nama = rapikan(re.sub(r"\*+$", "", b))
            n = seen[tahap_final][nama.casefold()] = seen[tahap_final].get(nama.casefold(), 0) + 1
            if n > 1:
                diganti.append(f"{tahap_final}: bahan '{nama}' muncul lagi, dinamai '{nama} ({n})'.")
                nama = f"{nama} ({n})"
            data[tahap_final].append({"nama": nama, "persen": persen, "gid": grup_d.get(r),
                                      "waktu": d, "ket": e,
                                      "bintang": bintang.group(0) if bintang else ""})
            terakhir = ("bahan", tahap_final)
            continue

        # ---- baris tanpa bahan kimia ----
        if a:
            berikut = baris[i + 1] if i + 1 < len(baris) else None
            if berikut and berikut["b"] and not berikut["a"] and berikut["c"] is not None:
                tahap_raw, tahap_final = a, None  # judul tahap; bahan di bawahnya miliknya
                if nilai_teks(d) or nilai_teks(e):
                    catatan.append(f"Baris {r}: waktu/keterangan pada judul '{a}' tidak diimpor.")
                continue
            nama = nama_unik(a)  # tahapan tanpa bahan kimia
            tanpa[nama] = {"w": [t for t in [_bersih(nilai_teks(d))] if t],
                           "k": [t for t in [_bersih(nilai_teks(e))] if t]}
            tahap_raw, tahap_final, terakhir = None, None, ("tanpa", nama)
            continue
        # lanjutan: tanpa nama tahap & bahan -> menambah waktu/keterangan langkah sebelumnya
        td, te = _bersih(nilai_teks(d)), _bersih(nilai_teks(e))
        if terakhir is None:
            catatan.append(f"Baris {r} dilewati (belum ada tahap): {td} {te}".strip())
        elif terakhir[0] == "tanpa":
            t = tanpa[terakhir[1]]
            if td and td not in t["w"]:
                t["w"].append(td)
            if te and te not in t["k"]:
                t["k"].append(te)
        else:
            tambahan = lanjutan.setdefault(terakhir[1], [])
            for tt in (f"{td} menit" if td else "", te):
                if tt and tt not in tambahan:
                    tambahan.append(tt)

    rec, info_fase, daftar_tanpa = [], {}, []
    for tahap, rows in data.items():
        fases = _bentuk_fase(rows)
        for no, fase in enumerate(fases, 1):
            waktu = next((nilai_teks(x["waktu"]) for x in fase if nilai_teks(x["waktu"])), "")
            kets = []
            for x in fase:
                for t in (_bersih(nilai_teks(x["ket"])),
                          f"{x['nama']}: {kaki[x['bintang']]}" if x["bintang"] in kaki else ""):
                    if t and t not in kets:
                        kets.append(t)
            if no == len(fases):  # baris lanjutan menempel di fase terakhir tahap ini
                kets += [t for t in lanjutan.get(tahap, []) if t not in kets]
            for x in fase:
                rec.append({"jenis_kulit": jenis, "tahapan": tahap, "senyawa": x["nama"],
                            "persen": x["persen"], "fase": no,
                            "satuan": "mL" if x["nama"].casefold().startswith("air") else "g"})
            if waktu or kets:
                info_fase[(jenis, tahap, no)] = (waktu, "; ".join(kets))
    for nama, t in tanpa.items():
        w, k = "; ".join(t["w"]), "; ".join(t["k"])
        daftar_tanpa.append({"jenis_kulit": jenis, "tahapan": nama, "waktu": w, "keterangan": k})
        info_fase[(jenis, nama, 1)] = (w, k)
    if len(catatan) > 12:
        catatan = catatan[:12] + [f"... dan {len(catatan) - 12} baris lain dilewati."]
    return {"formulasi": rec, "info": info_fase, "catatan": diganti + catatan,
            "format": "resep", "tanpa_bahan": daftar_tanpa}


def bagian_impor():
    """Impor formulasi dari file Kartu Kendali Proses atau file Resep (.xlsx)."""
    n = st.session_state.get("imp_n", 0)
    with st.expander("📥 Impor dari file Excel (Kartu Kendali Proses / Resep)",
                     expanded="impor" in st.session_state):
        st.caption("Format dikenali otomatis: **Kartu Kendali Proses** (kolom Tanggal | Proses | "
                   "Bahan Kimia | ...) atau **Resep** (kolom Tahap | Bahan | Jumlah (%) | Waktu | "
                   "Keterangan). Sel Waktu yang digabung dianggap satu fase. Tahapan tanpa bahan "
                   "kimia (mis. Aging, Shaving) ikut terdeteksi, hanya berisi waktu & keterangan. "
                   "Hasilnya bisa dicek dulu sebelum disimpan.")
        f = st.file_uploader("File Excel (.xlsx)", type=["xlsx"], key=f"imp_{n}")
        sheet, fmt, jenis_manual = None, None, ""
        if f is not None:
            try:
                deteksi = deteksi_format(f)
            except Exception as e:  # noqa: BLE001
                deteksi = {}
                st.error(f"❌ Gagal membuka file: {str(e)[:200]}")
            if not deteksi:
                st.warning("Format file tidak dikenali. Header tabel harus 'Tanggal | Proses' "
                           "(Kartu Kendali) atau 'Tahap | Bahan' (Resep).")
            else:
                sheet = (st.selectbox("Sheet", list(deteksi), key=f"imp_sheet_{n}")
                         if len(deteksi) > 1 else next(iter(deteksi)))
                fmt = deteksi[sheet]
                st.caption("Format terdeteksi: " + ("Kartu Kendali Proses" if fmt == "kartu"
                                                   else "Resep (Tahap | Bahan | Jumlah)"))
                if fmt == "resep":
                    jenis_manual = rapikan(st.text_input(
                        "Jenis kulit untuk resep ini (tidak ada di file)", key=f"imp_jenis_{n}",
                        placeholder="contoh: Biawak"))
        if fmt and st.button("🔍 Baca file"):
            st.session_state["imp_ver"] = st.session_state.get("imp_ver", 0) + 1
            try:
                if fmt == "resep":
                    if not jenis_manual:
                        raise ValueError("Isi dulu jenis kulit untuk resep ini.")
                    st.session_state["impor"] = baca_resep(f, sheet, jenis_manual.title())
                else:
                    st.session_state["impor"] = baca_formulasi(f, sheet)
            except Exception as e:  # noqa: BLE001
                st.session_state.pop("impor", None)
                st.error(f"❌ Gagal membaca file: {str(e)[:250]}")

        hasil = st.session_state.get("impor")
        if not hasil:
            return
        formulasi, tanpa = hasil["formulasi"], hasil.get("tanpa_bahan", [])
        if not formulasi and not tanpa:
            st.warning("Tidak ada tahapan atau bahan kimia yang terdeteksi.")
            return
        n_tahap = (pd.DataFrame(formulasi).groupby(["jenis_kulit", "tahapan"]).ngroups
                   if formulasi else 0)
        st.success(f"Terdeteksi {len(formulasi)} bahan kimia di {n_tahap} tahapan, "
                   f"dan {len(tanpa)} tahapan tanpa bahan kimia.")
        if hasil.get("format") == "resep" and formulasi:
            st.info("File resep tidak punya satuan. Default: Air = mL, bahan lain = g. "
                    "Ubah langsung di kolom Satuan pada tabel di bawah kalau perlu.")
        for c in hasil["catatan"]:
            st.warning(c)

        satuan_final = []
        if formulasi:
            prev = pd.DataFrame(formulasi)
            wk = [hasil["info"].get((r.jenis_kulit, r.tahapan, r.fase), ("", ""))
                  for r in prev.itertuples()]
            prev["waktu"], prev["keterangan"] = [w[0] for w in wk], [w[1] for w in wk]
            tabel = prev.rename(columns={
                "jenis_kulit": "Jenis Kulit", "tahapan": "Tahapan", "senyawa": "Senyawa",
                "persen": "Persen (%)", "satuan": "Satuan", "fase": "Fase",
                "waktu": "Waktu (menit)", "keterangan": "Keterangan"})
            edit = st.data_editor(
                tabel, hide_index=True, width="stretch",
                disabled=[c for c in tabel.columns if c != "Satuan"],
                column_config={"Satuan": st.column_config.SelectboxColumn(
                    "Satuan", options=list(SATUAN), required=True)},
                key=f"imp_editor_{n}_{st.session_state.get('imp_ver', 0)}")
            satuan_final = list(edit["Satuan"])
        if tanpa:
            st.markdown("**Tahapan tanpa bahan kimia** (hanya waktu & keterangan):")
            st.dataframe(pd.DataFrame(tanpa).rename(columns={
                "jenis_kulit": "Jenis Kulit", "tahapan": "Tahapan",
                "waktu": "Waktu (menit)", "keterangan": "Keterangan"}),
                hide_index=True, width="stretch")

        mode = st.radio("Kalau jenis kulit + tahapan + senyawa sudah ada di database:",
                        [MODE_LEWATI, MODE_TIMPA], key="imp_mode")
        c1, c2 = st.columns(2)
        if c1.button("💾 Simpan ke database", type="primary"):
            final = {**hasil, "formulasi": [
                {**r, "satuan": sat} for r, sat in zip(formulasi, satuan_final)]}
            try:
                pesan = simpan_impor(final, mode)
            except DBError as e:
                st.error(f"❌ {e}")
            else:
                st.session_state.pop("impor", None)
                st.session_state["imp_n"] = n + 1
                flash("warning" if pesan.startswith("⚠️") else "success", pesan)
        if c2.button("Batal"):
            st.session_state.pop("impor", None)
            st.session_state["imp_n"] = n + 1
            st.rerun()


def semua_tahapan(df: pd.DataFrame, info: dict, jenis) -> list:
    """Tahapan untuk satu jenis kulit: yang punya bahan kimia (urut input), lalu yang tanpa bahan."""
    ada = list(dict.fromkeys(df[df["jenis_kulit"] == jenis]["tahapan"])) if not df.empty else []
    tanpa = [t for (j, t, _) in info if j == jenis and t not in ada]
    return ada + list(dict.fromkeys(tanpa))


def bagian_info_tahapan(df: pd.DataFrame):
    """Atur waktu & keterangan per fase, atau tahapan tanpa bahan kimia."""
    st.divider()
    st.subheader("Waktu & keterangan per fase / tahapan tanpa bahan kimia")
    st.caption("Senyawa dengan nomor fase sama dikerjakan bersamaan (berbagi waktu & keterangan). "
               "Contoh Netralisasi: Air + Natrium Format = fase 1 (30 menit); "
               "Soda Kue = fase 2 (3 x 20, cek pH). Untuk tahapan yang tidak memakai bahan kimia "
               "(mis. Aging, Shaving), pilih '➕ Tahapan baru tanpa bahan kimia', lalu isi nama, "
               "waktu, dan keterangan apa yang dilakukan.")
    try:
        info = ambil_info()
    except DBError as e:
        st.error(f"❌ {e}")
        return
    jenis_opsi = sorted(set(df["jenis_kulit"]) | {j for (j, _, _) in info})
    if not jenis_opsi:
        st.info("Tambahkan formulasi dulu.")
        return

    c1, c2, c3 = st.columns([2, 2, 1])
    with c1:
        jenis = st.selectbox("Jenis kulit", jenis_opsi, key="info_jenis")
    with c2:
        pilih = st.selectbox("Tahapan", semua_tahapan(df, info, jenis) + [BARU_TAHAP],
                             key="info_tahap")
    if pilih == BARU_TAHAP:
        tahapan = rapikan(st.text_input("Nama tahapan baru (tanpa bahan kimia)",
                                        key="info_tahap_baru", placeholder="contoh: Aging")).title()
        if not tahapan:
            st.caption("Isi nama tahapan baru.")
            return
        fase_opsi = [1]
    else:
        tahapan = pilih
        sub_fase = df[(df["jenis_kulit"] == jenis) & (df["tahapan"] == tahapan)]["fase"].tolist()
        fase_opsi = sorted(set(sub_fase) | {f for (j, t, f) in info if j == jenis and t == tahapan})
    with c3:
        fase = int(st.selectbox("Fase", fase_opsi, key="info_fase"))

    sub = df[(df["jenis_kulit"] == jenis) & (df["tahapan"] == tahapan) & (df["fase"] == fase)]
    st.caption(f"Senyawa di fase ini: {', '.join(sub['senyawa'])}" if not sub.empty
               else "Tahapan tanpa bahan kimia (hanya waktu & keterangan).")

    lama_w, lama_k = info.get((jenis, tahapan, fase), ("", ""))
    kunci = f"{jenis}|{tahapan}|{fase}"  # key berubah -> isian ikut ter-reset ke data tersimpan
    waktu = st.text_input("Waktu (menit)", value=lama_w, key=f"waktu_{kunci}",
                          placeholder="contoh: 30  atau  3 x 20")
    ket = st.text_area("Keterangan (opsional)", value=lama_k, key=f"ket_{kunci}",
                       placeholder="contoh: pH minimum 4, uji tembus")
    if st.button("💾 Simpan waktu & keterangan"):
        waktu, ket = rapikan(waktu), (ket or "").strip()
        try:
            if not waktu and not ket and pilih != BARU_TAHAP:
                jalankan(lambda: db().table("tahapan_info").delete()
                         .eq("jenis_kulit", jenis).eq("tahapan", tahapan)
                         .eq("fase", fase).execute())
                flash("success", f"🗑️ Info {jenis} / {tahapan} / fase {fase} dikosongkan.")
            else:
                jalankan(lambda: db().table("tahapan_info").upsert(
                    {"jenis_kulit": jenis, "tahapan": tahapan, "fase": fase,
                     "waktu": waktu, "keterangan": ket},
                    on_conflict="jenis_kulit,tahapan,fase").execute())
                flash("success", f"✅ Tersimpan: {jenis} / {tahapan} / fase {fase}")
        except DBError as e:
            st.error(f"❌ {e}")

    if info:
        st.caption("Info yang sudah tersimpan:")
        ada = set(zip(df["jenis_kulit"], df["tahapan"]))
        st.dataframe(pd.DataFrame(
            [{"Jenis Kulit": j, "Tahapan": t, "Fase": f,
              "Bahan kimia": "ya" if (j, t) in ada else "tanpa bahan",
              "Waktu (menit)": w, "Keterangan": k}
             for (j, t, f), (w, k) in sorted(info.items())]),
            hide_index=True, width="stretch")


def halaman_developer():
    st.title("Developer: Kelola Formulasi")
    tampil_flash()
    try:
        df = ambil_semua()
    except DBError as e:
        st.error(f"❌ {e}")
        return

    jenis_list = sorted(df["jenis_kulit"].unique().tolist())
    tahap_list = list(dict.fromkeys(df["tahapan"].tolist()))

    bagian_impor()

    st.subheader("Tambah formulasi")
    c1, c2 = st.columns(2)
    with c1:
        jenis = pilih_atau_baru("Jenis kulit", jenis_list, "jenis")
    with c2:
        tahapan = pilih_atau_baru("Tahapan", tahap_list, "tahap")
    c3, c4, c5, c6 = st.columns([2, 2, 1, 1])
    with c3:
        senyawa = st.text_input("Nama senyawa kimia", key="senyawa")
    with c4:
        persen = st.number_input("Persentase (% dari berat kulit)",
                                 min_value=0.0, value=None, format="%.4f", key="persen")
    with c5:
        satuan = st.selectbox("Satuan output", list(SATUAN), key="satuan")
    with c6:
        fase = st.number_input("Fase", min_value=1, value=1, step=1, key="fase",
                               help="Senyawa dengan fase sama dikerjakan bersamaan "
                                    "(waktu & keterangan sama). Fase 1 dulu, lalu fase 2, dst.")

    if st.button("➕ Tambah", type="primary"):
        jenis, tahapan = rapikan(jenis).title(), rapikan(tahapan).title()
        senyawa = rapikan(senyawa)
        if not (jenis and tahapan and senyawa):
            st.warning("⚠️ Jenis kulit, tahapan, dan senyawa wajib diisi.")
        elif persen is None or persen <= 0:
            st.warning("⚠️ Persentase harus angka lebih dari 0.")
        else:
            try:
                jalankan(lambda: db().table(TABEL).insert({
                    "jenis_kulit": jenis, "tahapan": tahapan,
                    "senyawa": senyawa, "persen": persen,
                    "satuan": satuan, "fase": int(fase)}).execute())
            except DBError as e:
                st.error(f"❌ {e}")
            else:
                flash("success", f"✅ Ditambahkan: {jenis} / {tahapan} / {senyawa} ({persen}%, output {satuan}, fase {int(fase)})")

    bagian_info_tahapan(df)

    st.divider()
    st.subheader("Data formulasi (klik satu baris untuk edit / hapus)")
    if df.empty:
        st.info("Belum ada data.")
        return

    event = st.dataframe(df, hide_index=True, width="stretch",
                         on_select="rerun", selection_mode="single-row")
    baris = event.selection.rows
    if not baris:
        return

    row = df.iloc[baris[0]]
    st.write(f"Terpilih: **{row['jenis_kulit']} / {row['tahapan']} / {row['senyawa']}** "
             f"(ID {int(row['id'])}, {row['persen']}%)")
    ce1, ce2, ce3 = st.columns([2, 1, 1])
    with ce1:
        persen_baru = st.number_input("Persentase baru (%)", min_value=0.0,
                                      value=float(row["persen"]), format="%.4f",
                                      key=f"persen_baru_{int(row['id'])}")
    with ce2:
        satuan_baru = st.selectbox("Satuan baru", list(SATUAN),
                                   index=list(SATUAN).index(row["satuan"]) if row["satuan"] in SATUAN else 0,
                                   key=f"satuan_baru_{int(row['id'])}")
    with ce3:
        fase_baru = st.number_input("Fase baru", min_value=1, value=int(row["fase"]), step=1,
                                    key=f"fase_baru_{int(row['id'])}")
    cu, ch = st.columns(2)
    if cu.button("✏️ Ubah persentase, satuan & fase"):
        if persen_baru <= 0:
            st.warning("⚠️ Persentase harus angka lebih dari 0.")
        else:
            try:
                jalankan(lambda: db().table(TABEL).update({"persen": persen_baru, "satuan": satuan_baru,
                                                              "fase": int(fase_baru)})
                         .eq("id", int(row["id"])).execute())
            except DBError as e:
                st.error(f"❌ {e}")
            else:
                flash("success", f"✅ ID {int(row['id'])} diubah: {persen_baru}%, satuan {satuan_baru}, fase {int(fase_baru)}")
    if ch.button("🗑️ Hapus baris"):
        try:
            jalankan(lambda: db().table(TABEL).delete().eq("id", int(row["id"])).execute())
        except DBError as e:
            st.error(f"❌ {e}")
        else:
            flash("success", f"🗑️ Baris ID {int(row['id'])} dihapus.")


# ---------------------------------------------------------------------------
# Interface User
# ---------------------------------------------------------------------------
UNIT_XLS = {"g": "gr", "kg": "kg", "mL": "ml", "L": "L"}          # tulisan di Excel
UNIT_BALIK = {"gr": "g", "g": "g", "kg": "kg", "ml": "mL", "l": "L"}  # baca dari Excel
POLA_BERAT = re.compile(
    r"^Berat Kulit\s+(.+?)\s+(?:Kering|Setelah\b.*?)\s*:\s*([\d.,]+)\s*(?:gr|g)\b", re.I)
POLA_JUMLAH = re.compile(
    r"^\s*([\d.,]+)\s*\(\s*([\d.,]+)\s*(gr|g|kg|ml|l)\s*\)\s*$", re.I)


def tanggal_hari_ini() -> str:
    return datetime.now(timezone(timedelta(hours=7))).strftime("%Y-%m-%d")  # WIB


def fmt_angka(x, maks: int = 2, tetap: bool = False) -> str:
    """Format angka gaya Indonesia (koma desimal). tetap=True -> selalu `maks` desimal."""
    if x is None or pd.isna(x):
        return ""
    d = Decimal(str(round(float(x), 6))).quantize(Decimal(1).scaleb(-maks), ROUND_HALF_UP)
    teks = f"{d:.{maks}f}" if tetap else format(d.normalize(), "f")
    return teks.replace(".", ",")


def angka(teks: str) -> float:
    return float(str(teks).strip().replace(",", "."))


def hitung_langkah(df: pd.DataFrame, jenis, tahapan, berat, info=None) -> pd.DataFrame:
    info = info or {}
    h = df[(df["jenis_kulit"] == jenis) & (df["tahapan"] == tahapan)].copy()
    h = h.sort_values(["fase", "id"], kind="stable")
    # berat kulit dalam g -> jumlah senyawa dikonversi ke satuan masing-masing
    h["jumlah"] = (h["persen"] / 100 * berat * h["satuan"].map(SATUAN).fillna(1)).round(4)
    h["waktu"] = [info.get((jenis, tahapan, f), ("", ""))[0] for f in h["fase"]]
    h["keterangan"] = [info.get((jenis, tahapan, f), ("", ""))[1] for f in h["fase"]]
    return h[["senyawa", "persen", "jumlah", "satuan", "fase", "waktu", "keterangan"]]


def ringkas_satuan(rows: pd.DataFrame) -> str:
    g = rows.groupby("satuan")["jumlah"].sum()
    return " + ".join(f"{fmt_angka(v, 2, True)} {UNIT_XLS.get(k, k)}" for k, v in g.items())


def tabel_tampil(rows) -> pd.DataFrame:
    kolom = ["senyawa", "persen", "jumlah", "satuan", "fase", "waktu", "keterangan"]
    d = pd.DataFrame(rows).reindex(columns=kolom)
    d[["waktu", "keterangan"]] = d[["waktu", "keterangan"]].fillna("")
    return d.rename(columns={
        "senyawa": "Senyawa", "persen": "Persen (%)", "jumlah": "Jumlah", "satuan": "Satuan",
        "fase": "Fase", "waktu": "Waktu (menit)", "keterangan": "Keterangan"})


KONFIG_TABEL = {"Jumlah": st.column_config.NumberColumn("Jumlah", format="%.2f")}


def nilai_teks(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return fmt_angka(v)
    return str(v).strip()


def ubah_tanggal(a) -> str:
    if isinstance(a, date):  # termasuk datetime
        return a.strftime("%Y-%m-%d")
    if isinstance(a, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(a.strip(), fmt).strftime("%Y-%m-%d")
            except ValueError:
                pass
    return ""


def baca_file(f, sheet=None):
    """Baca file Kartu Kendali Proses (hasil export app ini atau tulisan manual).
    Return (daftar_langkah, info) dengan info = {'penyamak': ...}.
    Baris yang kolom Waktu-nya digabung (merge) dianggap satu fase."""
    wb = load_workbook(f, data_only=True)
    if sheet:
        ws = wb[sheet]
    else:
        ws = wb["Kartu Kendali"] if "Kartu Kendali" in wb.sheetnames else wb.worksheets[0]

    info, mulai = {"penyamak": ""}, None
    for r in range(1, min(ws.max_row, 40) + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str):
            kunci = a.strip().lower()
            if kunci == "penyamak" and isinstance(b, str):
                info["penyamak"] = b.lstrip(": ").strip()
            if kunci == "tanggal" and isinstance(b, str) and b.strip().lower() == "proses":
                mulai = r + 1
                break
    if mulai is None:
        raise ValueError("Header tabel (Tanggal, Proses, Bahan Kimia, ...) tidak ditemukan. "
                         "Pakai file Kartu Kendali Proses.")

    grup_e = {}  # baris -> baris pertama dari sel Waktu yang digabung
    for rng in ws.merged_cells.ranges:
        if rng.min_col <= 5 <= rng.max_col:
            for rr in range(rng.min_row, rng.max_row + 1):
                grup_e[rr] = rng.min_row

    daftar, cur, jenis, berat, tgl, gagal = [], None, None, None, "", 0
    for r in range(mulai, ws.max_row + 1):
        a, b, c, d = (ws.cell(r, k).value for k in (1, 2, 3, 4))
        if a not in (None, ""):
            tgl = ubah_tanggal(a) or tgl
        if isinstance(b, str) and b.strip():
            m = POLA_BERAT.match(b.strip())
            if m:  # baris "Berat Kulit ... : N gr"
                jenis, berat, cur = m.group(1).strip(), angka(m.group(2)), None
                continue
            if jenis is None:
                continue
            cur = {"tanggal": tgl, "jenis": jenis, "tahapan": b.strip(),
                   "berat": berat, "rows": [], "_grup": {}}
            daftar.append(cur)
        if cur is None:
            continue
        gid = grup_e.get(r, -r)  # baris tak digabung = fase sendiri
        g = cur["_grup"].setdefault(gid, {"w": [], "k": []})
        for v, kunci in ((ws.cell(r, 5).value, "w"), (ws.cell(r, 6).value, "k")):
            t = nilai_teks(v)
            if t and t not in g[kunci]:
                g[kunci].append(t)
        if isinstance(c, str) and c.strip() and d is not None:
            m = POLA_JUMLAH.match(str(d))
            if m:
                cur["rows"].append({
                    "senyawa": c.strip(), "persen": angka(m.group(1)),
                    "jumlah": angka(m.group(2)),
                    "satuan": UNIT_BALIK[m.group(3).lower()], "_gid": gid})
            else:
                gagal += 1

    for e in daftar:
        urut = {}
        for row in e["rows"]:
            gid = row.pop("_gid")
            urut.setdefault(gid, len(urut) + 1)
            row["fase"] = urut[gid]
            row["waktu"] = "; ".join(e["_grup"][gid]["w"])
            row["keterangan"] = "; ".join(e["_grup"][gid]["k"])
        if not e["rows"]:  # tahapan tanpa bahan kimia: hanya waktu & keterangan
            w = [t for g in e["_grup"].values() for t in g["w"]]
            k = [t for g in e["_grup"].values() for t in g["k"]]
            e["tanpa_bahan"] = {"waktu": "; ".join(dict.fromkeys(w)),
                                "keterangan": "; ".join(dict.fromkeys(k))}
        e.pop("_grup")
    daftar = [e for e in daftar if e["rows"] or "tanpa_bahan" in e]
    if not daftar:
        raise ValueError("Tidak ada langkah yang bisa dibaca dari file ini.")
    info["gagal"] = gagal
    return daftar, info


def buat_excel(daftar: list, penyamak: str = "") -> bytes:
    """Kartu Kendali Proses: 1 sheet kartu + 1 sheet ringkasan total per senyawa."""
    font, tebal = Font(name="Arial", size=11), Font(name="Arial", size=11, bold=True)
    sisi = Side(style="thin")
    kotak = Border(left=sisi, right=sisi, top=sisi, bottom=sisi)
    tengah = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Kartu Kendali"
    for kol, lebar in zip("ABCDEF", (13, 22, 38, 22, 16, 26)):
        ws.column_dimensions[kol].width = lebar

    ws.merge_cells("A1:F1")
    ws["A1"] = "KARTU KENDALI PROSES"
    ws["A1"].font, ws["A1"].alignment = tebal, Alignment(horizontal="center")

    jenis_unik = list(dict.fromkeys(e["jenis"] for e in daftar))
    for baris, (label, isi) in zip((3, 4, 5), (
            ("Jenis Kulit", ": Kulit " + ", ".join(jenis_unik)),
            ("Penyamak", ": " + (penyamak or "").strip()),
            ("No. order", ":"))):
        ws.cell(baris, 1, label).font = font
        ws.cell(baris, 2, isi).font = font

    for k, judul in enumerate(("Tanggal", "Proses", "Bahan Kimia", "Jumlah (%)",
                               "Waktu (menit)", "Keterangan"), 1):
        c = ws.cell(7, k, judul)
        c.font, c.alignment, c.border = tebal, tengah, kotak

    def isi_baris(r, nilai):
        for k in range(1, 7):
            c = ws.cell(r, k)
            c.font, c.alignment, c.border = font, tengah, kotak
            if k in nilai:
                c.value = nilai[k]

    def tulis_waktu_ket(awal, akhir, waktu, ket):
        """Isi kolom Waktu & Keterangan, digabung vertikal kalau lebih dari satu baris."""
        waktu, ket = (waktu or "").strip(), (ket or "").strip()
        if waktu:
            ws.cell(awal, 5).value = (angka(waktu) if re.fullmatch(r"\d+([.,]\d+)?", waktu)
                                      else waktu)
        if ket:
            ws.cell(awal, 6).value = ket
        if akhir > awal:
            for kol in (5, 6):
                ws.merge_cells(start_row=awal, start_column=kol, end_row=akhir, end_column=kol)
        # tinggi baris agar keterangan panjang tidak terpotong (sel gabungan tidak auto-fit)
        perlu = sum(max(1, -(-len(x) // 24)) for x in ket.split("\n")) if ket else 1
        jumlah_baris = akhir - awal + 1
        if perlu > jumlah_baris:
            ws.row_dimensions[awal].height = 15.75 * (1 + perlu - jumlah_baris)

    r, tgl_sebelumnya = 8, None
    for i, e in enumerate(daftar):
        awal_langkah = r
        if not e["rows"]:
            # Tahapan tanpa bahan kimia: hanya nama proses, waktu, dan keterangan
            tb = e.get("tanpa_bahan", {})
            isi_baris(r, {2: e["tahapan"]})
            tulis_waktu_ket(r, r, tb.get("waktu"), tb.get("keterangan"))
            r += 1
        else:
            # Baris berat kulit (digabung B:F)
            kondisi = "Kering" if i == 0 else f"Setelah {daftar[i - 1]['tahapan']}"
            isi_baris(r, {2: f"Berat Kulit {e['jenis']} {kondisi}: {fmt_angka(e['berat'])} gr"})
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            r += 1
            # Bahan kimia per fase; Waktu & Keterangan digabung per fase
            pertama = True
            for _, grup in groupby(e["rows"], key=lambda x: x.get("fase", 1)):
                grup = list(grup)
                awal = r
                for row in grup:
                    isi_baris(r, {
                        2: e["tahapan"] if pertama else None,
                        3: row["senyawa"],
                        4: f"{fmt_angka(row['persen'], 4)} "
                           f"({fmt_angka(row['jumlah'], 2, True)} "
                           f"{UNIT_XLS.get(row['satuan'], row['satuan'])})"})
                    pertama = False
                    r += 1
                tulis_waktu_ket(awal, r - 1, grup[0].get("waktu"), grup[0].get("keterangan"))
        # Tanggal otomatis di baris pertama langkah, hanya saat tanggal berganti
        if e["tanggal"] and e["tanggal"] != tgl_sebelumnya:
            try:
                ws.cell(awal_langkah, 1).value = datetime.strptime(e["tanggal"], "%Y-%m-%d").date()
                ws.cell(awal_langkah, 1).number_format = "dd/mm/yyyy"
            except ValueError:
                ws.cell(awal_langkah, 1).value = e["tanggal"]
        tgl_sebelumnya = e["tanggal"]

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Sheet 2: total per senyawa dari semua langkah
    rs = wb.create_sheet("Ringkasan")
    semua = pd.DataFrame([x for e in daftar for x in e["rows"]],
                         columns=["senyawa", "persen", "jumlah", "satuan", "fase",
                                  "waktu", "keterangan"])
    total = semua.groupby(["senyawa", "satuan"], as_index=False)["jumlah"].sum()
    for k, judul in enumerate(("Senyawa", "Total Jumlah", "Satuan"), 1):
        c = rs.cell(1, k, judul)
        c.font, c.alignment, c.border = tebal, tengah, kotak
    for n, row in enumerate(total.itertuples(index=False), 2):
        for k, v in enumerate((row.senyawa, round(float(row.jumlah), 2),
                               UNIT_XLS.get(row.satuan, row.satuan)), 1):
            c = rs.cell(n, k, v)
            c.font, c.alignment, c.border = font, tengah, kotak
        rs.cell(n, 2).number_format = "0.00"
    for kol, lebar in zip("ABC", (32, 16, 10)):
        rs.column_dimensions[kol].width = lebar

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def halaman_user():
    st.title("Hitung Kebutuhan Bahan Kimia")
    tampil_flash()
    if "daftar" not in st.session_state:
        st.session_state["daftar"] = []
    daftar = st.session_state["daftar"]

    try:
        df = ambil_semua()
    except DBError as e:
        st.error(f"❌ {e}")
        return

    try:
        info = ambil_info()
    except DBError as e:
        info = {}
        st.warning(f"Waktu & keterangan tahapan tidak bisa dimuat: {e}")

    # --- Lanjutkan dari file sebelumnya ---
    with st.expander("📂 Lanjutkan dari file sebelumnya"):
        st.caption("Upload file Kartu Kendali Proses (.xlsx) hasil export sebelumnya. "
                   "Langkah di dalamnya dimuat ke daftar, lalu kamu bisa menambah langkah baru "
                   "dan export ulang.")
        f = st.file_uploader("File Excel (.xlsx)", type=["xlsx"],
                             key=f"upl_{st.session_state.get('upl_n', 0)}")
        if f is not None and st.button("Muat ke daftar"):
            try:
                baru, info = baca_file(f)
            except Exception as e:  # noqa: BLE001
                st.error(f"❌ Gagal membaca file: {str(e)[:250]}")
            else:
                daftar[:0] = baru  # langkah lama ditaruh di depan
                if info["penyamak"]:
                    st.session_state["penyamak"] = info["penyamak"]
                st.session_state["upl_n"] = st.session_state.get("upl_n", 0) + 1
                pesan = f"✅ {len(baru)} langkah dari file dimuat ke daftar."
                if info["gagal"]:
                    pesan += f" ({info['gagal']} baris bahan kimia tidak terbaca.)"
                flash("success", pesan)

    penyamak = st.text_input("Penyamak (contoh: PKL UGM, magang, dll)", key="penyamak")

    # --- Hitung satu langkah ---
    st.subheader("Tambah langkah")
    jenis_opsi = sorted(set(df["jenis_kulit"]) | {j for (j, _, _) in info})
    if not jenis_opsi:
        st.info("Belum ada formulasi. Hubungi developer.")
    else:
        c1, c2, c3 = st.columns(3)
        with c1:
            jenis = st.selectbox("Jenis kulit", jenis_opsi, index=None,
                                 placeholder="Pilih jenis kulit")
        with c2:
            tahap_opsi = semua_tahapan(df, info, jenis) if jenis else []
            tahapan = st.selectbox("Tahapan", tahap_opsi, index=None,
                                   placeholder="Pilih tahapan")
        with c3:
            berat = st.number_input("Berat kulit (g)", min_value=0.0, value=None,
                                    format="%.2f")

        if jenis and tahapan:
            ada_bahan = not df[(df["jenis_kulit"] == jenis) & (df["tahapan"] == tahapan)].empty
            if not ada_bahan:
                # Tahapan tanpa bahan kimia: berat kulit tidak diperlukan
                w, k = info.get((jenis, tahapan, 1), ("", ""))
                st.info(f"Tahapan tanpa bahan kimia · Waktu: {w or '-'} menit · "
                        f"Keterangan: {k or '-'}")
                if st.button("➕ Tambah ke daftar", type="primary"):
                    daftar.append({"tanggal": tanggal_hari_ini(), "jenis": jenis,
                                   "tahapan": tahapan, "berat": None, "rows": [],
                                   "tanpa_bahan": {"waktu": w, "keterangan": k}})
                    flash("success", f"✅ Langkah {len(daftar)} ditambahkan: {jenis} / {tahapan}")
            elif berat is None or berat <= 0:
                st.info("Masukkan berat kulit (g) lebih dari 0.")
            else:
                rows = hitung_langkah(df, jenis, tahapan, berat, info)
                st.dataframe(tabel_tampil(rows), hide_index=True,
                             width="stretch", column_config=KONFIG_TABEL)
                st.markdown(f"**Total:** {ringkas_satuan(rows)}")
                st.caption("ml dan L dihitung dari gram dengan asumsi massa jenis 1 g = 1 ml.")
                if st.button("➕ Tambah ke daftar", type="primary"):
                    daftar.append({"tanggal": tanggal_hari_ini(), "jenis": jenis,
                                   "tahapan": tahapan, "berat": float(berat),
                                   "rows": rows.to_dict("records")})
                    flash("success", f"✅ Langkah {len(daftar)} ditambahkan: {jenis} / {tahapan}")

    # --- Daftar langkah + export ---
    st.divider()
    st.subheader(f"Daftar langkah ({len(daftar)})")
    if not daftar:
        st.info("Belum ada langkah. Hitung lalu klik 'Tambah ke daftar'.")
        return

    for i, e in enumerate(daftar):
        judul = f"{i + 1}. {e['jenis']} / {e['tahapan']}"
        if e["rows"] and e.get("berat") is not None:
            judul += f" · {fmt_angka(e['berat'])} g"
        if e["tanggal"]:
            judul += f" · {e['tanggal']}"
        with st.expander(judul):
            if e["rows"]:
                st.dataframe(tabel_tampil(e["rows"]), hide_index=True,
                             width="stretch", column_config=KONFIG_TABEL)
            else:
                tb = e.get("tanpa_bahan", {})
                st.caption(f"Tanpa bahan kimia · Waktu: {tb.get('waktu') or '-'} menit · "
                           f"Keterangan: {tb.get('keterangan') or '-'}")
            if st.button("🗑️ Hapus langkah ini", key=f"hapus_{i}"):
                daftar.pop(i)
                flash("success", "Langkah dihapus dari daftar.")

    st.download_button("📥 Export Kartu Kendali Proses (.xlsx)", buat_excel(daftar, penyamak),
                       file_name=f"kartu_kendali_proses_{tanggal_hari_ini()}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       type="primary")
    if st.button("Kosongkan daftar"):
        daftar.clear()
        flash("success", "Daftar dikosongkan.")


# ---------------------------------------------------------------------------
# Router
# ---------------------------------------------------------------------------
role = st.session_state.get("role")
if role is None:
    halaman_login()
else:
    logout_sidebar()
    if role == "dev":
        halaman_developer()
    else:
        halaman_user()
